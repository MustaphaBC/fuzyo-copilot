"""Multi-format document loaders for Fuzyo RAG ingestion."""

from __future__ import annotations

import logging
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

_CHUNK_SIZE = 800
_CHUNK_OVERLAP = 100


@dataclass
class DocumentChunk:
    content: str
    metadata: dict = field(default_factory=dict)


def chunk_text(text: str, *, chunk_size: int = _CHUNK_SIZE, overlap: int = _CHUNK_OVERLAP) -> list[str]:
    """Paragraph-aware splitter with character overlap."""
    cleaned = re.sub(r"\n{3,}", "\n\n", text.strip())
    if not cleaned:
        return []

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", cleaned) if p.strip()]
    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= chunk_size:
            current = candidate
            continue
        if current:
            chunks.append(current)
            # Overlap from end of previous chunk
            current = current[-overlap:] if overlap and len(current) > overlap else ""
            candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= chunk_size:
            current = candidate
            continue
        # Hard-split long paragraph
        start = 0
        while start < len(candidate):
            end = min(start + chunk_size, len(candidate))
            piece = candidate[start:end].strip()
            if piece:
                chunks.append(piece)
            if end >= len(candidate):
                break
            start = max(end - overlap, start + 1)
        current = ""

    if current:
        chunks.append(current)
    return chunks


class BaseDocumentLoader(ABC):
    """Abstract loader: extract text, then chunk with metadata."""

    file_type: str = "unknown"

    @abstractmethod
    def extract_text(self, path: Path) -> str:
        """Extract raw text from a file. Raises on hard failure."""

    def load(self, path: Path | str, *, sdlc_phase: int = 1) -> list[DocumentChunk]:
        file_path = Path(path)
        try:
            text = self.extract_text(file_path)
        except Exception as exc:  # noqa: BLE001 — graceful parse fallback
            logger.warning("Failed to parse %s: %s", file_path, exc)
            return []

        pieces = chunk_text(text)
        return [
            DocumentChunk(
                content=piece,
                metadata={
                    "file_type": self.file_type,
                    "sdlc_phase": sdlc_phase,
                    "source": str(file_path),
                    "chunk_index": index,
                },
            )
            for index, piece in enumerate(pieces)
        ]


class MarkdownLoader(BaseDocumentLoader):
    file_type = "md"

    def extract_text(self, path: Path) -> str:
        return path.read_text(encoding="utf-8")


class PdfLoader(BaseDocumentLoader):
    file_type = "pdf"

    def extract_text(self, path: Path) -> str:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(part for part in pages if part.strip())


class DocxLoader(BaseDocumentLoader):
    file_type = "docx"

    def extract_text(self, path: Path) -> str:
        from docx import Document

        document = Document(str(path))
        return "\n\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())


class XlsxLoader(BaseDocumentLoader):
    file_type = "xlsx"

    def extract_text(self, path: Path) -> str:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        sections: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                sections.append(f"# Sheet: {sheet.title}\n" + "\n".join(rows))
        workbook.close()
        return "\n\n".join(sections)


class PlainTextLoader(BaseDocumentLoader):
    """UTF-8 text loader for source, config, and CSV files."""

    file_type = "text"

    def __init__(self, file_type: str = "text") -> None:
        self.file_type = file_type

    def extract_text(self, path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="ignore")


def _text_loader(file_type: str) -> type[BaseDocumentLoader]:
    class _Loader(PlainTextLoader):
        def __init__(self) -> None:
            super().__init__(file_type=file_type)

    _Loader.file_type = file_type
    return _Loader


_LOADERS: dict[str, type[BaseDocumentLoader]] = {
    ".md": MarkdownLoader,
    ".markdown": MarkdownLoader,
    ".pdf": PdfLoader,
    ".docx": DocxLoader,
    ".xlsx": XlsxLoader,
    ".csv": _text_loader("csv"),
    ".py": _text_loader("py"),
    ".ts": _text_loader("ts"),
    ".tsx": _text_loader("tsx"),
    ".js": _text_loader("js"),
    ".jsx": _text_loader("jsx"),
    ".go": _text_loader("go"),
    ".json": _text_loader("json"),
    ".yml": _text_loader("yml"),
    ".yaml": _text_loader("yaml"),
    ".toml": _text_loader("toml"),
    ".sql": _text_loader("sql"),
    ".txt": _text_loader("txt"),
}


def get_loader(path: Path | str) -> BaseDocumentLoader:
    """Return a loader instance for the file extension."""
    suffix = Path(path).suffix.lower()
    loader_cls = _LOADERS.get(suffix)
    if loader_cls is None:
        raise ValueError(f"Unsupported document type: {suffix or '(none)'}")
    return loader_cls()
